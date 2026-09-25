# -*- coding: utf-8 -*-
r"""TOCHKA - Руководство пользователя (RU). Генератор DOCX.

Стиль наследует _gen_manual_ru.py из STUKACH/LAMPOCHKA (Arial, синие заголовки,
фирменные таблицы). Запуск:  python _gen_manual_ru.py
Выход:   D:\AI\ZCode\Project\TOCHKA\docs\TOCHKA_Manual_RU.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r'D:\AI\ZCode\Project\TOCHKA\docs\TOCHKA_Manual_RU.docx'

doc = Document()

for section in doc.sections:
    section.top_margin    = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin   = Inches(1)
    section.right_margin  = Inches(1)

# ── базовые стили ──────────────────────────────────────────────────────────

normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.space_before = Pt(0)
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.find(qn('w:rFonts'))
rfonts.set(qn('w:cs'), 'Arial')

for lvl, size in (('Heading 1', 15), ('Heading 2', 12.5), ('Heading 3', 11)):
    st = doc.styles[lvl]
    st.font.name = 'Arial'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D) if lvl == 'Heading 1' else RGBColor(0x2E, 0x74, 0xB5)
    st.paragraph_format.space_before = Pt(14 if lvl == 'Heading 1' else 10)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.keep_with_next = True

# ── хелперы ────────────────────────────────────────────────────────────────

HDR_BG  = '2E75B5'
ALT_ROW = 'F5F8FB'
BLUE_H1 = RGBColor(0x1F, 0x49, 0x7D)
GREY    = RGBColor(0x55, 0x55, 0x55)


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color='CCCCCC'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcB = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)
        tcB.append(b)
    tcPr.append(tcB)


def cell_para(cell, text, bold=False, size=9, color=None, italic=False):
    p_ = cell.paragraphs[0]
    p_.paragraph_format.space_before = Pt(2)
    p_.paragraph_format.space_after = Pt(2)
    p_.paragraph_format.line_spacing = 1.1
    run = p_.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def table_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def add_table(doc, rows, col_widths_cm):
    table = doc.add_table(rows=0, cols=len(col_widths_cm))
    table.style = 'Table Grid'
    table.autofit = False
    table_margins(table)

    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        big = len(rows) > 2
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trPr.append(cant)
        if big and r_idx == 0:
            th = OxmlElement('w:tblHeader')
            trPr.append(th)

        glue = not big and r_idx < len(rows) - 1
        for c_idx, (text, w) in enumerate(zip(row_data, col_widths_cm)):
            cell = row.cells[c_idx]
            cell.width = Cm(w)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

            if r_idx == 0:
                set_cell_bg(cell, HDR_BG)
                cell_para(cell, text, bold=True, size=9, color=(0xFF, 0xFF, 0xFF))
            else:
                if r_idx % 2 == 0:
                    set_cell_bg(cell, ALT_ROW)
                if c_idx == 0:
                    cell_para(cell, text, bold=True, size=9)
                else:
                    cell_para(cell, text, size=9)

            if glue:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.line_spacing = 1.0
    return table


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def p(doc, text, italic=False, grey=False, bullet=False):
    par = doc.add_paragraph(style='List Bullet' if bullet else None)
    run = par.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(9.5)
    run.font.italic = italic
    if grey:
        run.font.color.rgb = GREY
    return par


def kv_note(doc, text):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = GREY


def mono(doc, text):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.bold = True
    par.paragraph_format.space_after = Pt(3)
    return par


def add_toc_field(doc):
    par = doc.add_paragraph()
    run = par.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = r'TOC \o "1-2" \h \z \u'
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = 'Оглавление: откройте документ в Word/LibreOffice и обновите поле (F9), чтобы заполнить номера страниц.'
    fld_sep.append(t)
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(fld_end)

    br = doc.add_paragraph()
    run_br = br.add_run()
    pb = OxmlElement('w:br')
    pb.set(qn('w:type'), 'page')
    run_br._r.append(pb)


def add_footer_pagenum(doc):
    footer = doc.sections[0].footer
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = r'PAGE \* arabic \* MERGEFORMAT'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(fld_end)
    run.font.name = 'Arial'
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


# ════════════════════════════════════════════════════════════════════════════
# ТИТУЛ
# ════════════════════════════════════════════════════════════════════════════

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(30)
title_p.paragraph_format.space_after = Pt(4)
r = title_p.add_run('TOCHKA')
r.font.name = 'Arial'; r.font.size = Pt(30); r.font.bold = True
r.font.color.rgb = BLUE_H1

title2 = doc.add_paragraph()
title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
title2.paragraph_format.space_after = Pt(10)
r = title2.add_run('Руководство пользователя')
r.font.name = 'Arial'; r.font.size = Pt(18); r.font.bold = True
r.font.color.rgb = BLUE_H1

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(4)
r = sub.add_run('Пивот-тулкит для Blender')
r.font.name = 'Arial'; r.font.size = Pt(11); r.font.color.rgb = GREY

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub2.paragraph_format.space_after = Pt(20)
r = sub2.add_run('Версия для Blender 4.2+ - v1.1.0')
r.font.name = 'Arial'; r.font.size = Pt(10.5); r.font.color.rgb = GREY

p(doc, 'TOCHKA - набор инструментов для работы с пивотами объектов: поставить пивот в центр '
       'выделения одним нажатием, перетащить его мышью по поверхности со снапом к вершинам, '
       'повернуть ориентацию, не двигая геометрию, выровнять по нормали грани или по ребру - '
       'и проверить пивоты сцены на соответствие конвенциям перед сдачей. Вместо плясок с '
       '3D-курсором вокруг штатного Set Origin - прямые визуальные операторы с живым превью.')

kv_note(doc, 'github.com/abyrvalg379/tochka')

# ════════════════════════════════════════════════════════════════════════════
# ОГЛАВЛЕНИЕ
# ════════════════════════════════════════════════════════════════════════════

h1(doc, 'Содержание')
add_toc_field(doc)

# ════════════════════════════════════════════════════════════════════════════
# 1. О ПРОГРАММЕ
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '1. О программе')

h2(doc, '1.1 Что делает TOCHKA')
p(doc, 'Ключевые возможности:', bullet=False)
for b in (
    'Origin to Selection - пивот в центр выделения клавишей D, с якорями Median / Bottom / Top;',
    'Drag Pivot - перетаскивание пивота мышью по любой видимой геометрии, со снапом к вершинам и осевыми констеинами;',
    'Rotate Pivot - поворот ориентации пивота без движения геометрии, для анимируемых частей;',
    'Align Pivot to Normal - ось Z пивота следует нормали грани под курсором;',
    'Align Pivot to Edge - ось X пивота ложится вдоль ребра;',
    'Pivot Audit - проверка пивотов выделенных мешей по конвенциям пайплайна с пакетным исправлением;',
    'честный undo: каждый инструмент коммитится одним шагом;',
    'сторож keymap: расширение само восстанавливает повреждённую системную раскладку при включении.',
):
    p(doc, b, bullet=True)

h2(doc, '1.2 Философия')
p(doc, 'Три принципа. Первый - пивот отдельный от геометрии: поворот и выравнивание ориентации '
       'меняют только матрицу объекта, вершины остаются на местах с точностью до машинного эпсилона. '
       'Второй - живое превью: каждый модальный инструмент показывает результат до подтверждения, '
       'и до коммита сцена не меняется - отмена (ПКМ или Esc) откатывает всё чисто. Третий - '
       'предсказуемый undo: любая операция коммитится ровно одним шагом истории.')

h2(doc, '1.3 Инструменты и клавиши')
add_table(doc, [
    ('Инструмент', 'Клавиша', 'Где ещё'),
    ('Origin to Selection', 'D · DD · DDD', 'панель, pie-меню'),
    ('Drag Pivot', 'Ctrl+D', 'панель, pie-меню'),
    ('Rotate Pivot', 'Ctrl+Alt+D (объектный режим)', 'панель, pie-меню'),
    ('Align Pivot to Normal', '-', 'панель, pie-меню'),
    ('Align Pivot to Edge', '-', 'панель, pie-меню'),
    ('Pivot Audit', '-', 'панель'),
], [5.2, 5.6, 4.4])
p(doc, 'Pie-меню со всеми операторами открывается кнопкой в N-панели.')

# ════════════════════════════════════════════════════════════════════════════
# 2. УСТАНОВКА
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '2. Установка')
for b in (
    'Скачайте zip последнего релиза: github.com/abyrvalg379/tochka → Releases → Latest → ассет tochka_v*.zip.',
    'Blender → Edit → Preferences → Get Extensions → ⚙ → Install from Disk → выберите zip. Так же возможна установка через drag-and-drop.',
    'Расширение включится само как TOCHKA; вкладка появится в N-панели 3D-вьюпорта (клавиша N). Версия показана прямо в заголовке панели.',
    'Обновление: установите новый zip тем же способом поверх старой версии.',
):
    p(doc, b, bullet=True)
p(doc, 'Протестировано в Blender 5.2. При каждом включении TOCHKA проверяет системную раскладку '
       '(keymap Object Mode и Mesh): если трансформ-бинды повреждены - кем бы то ни было - они '
       'восстанавливаются из заводских, в консоль печатается «TOCHKA: repaired damaged ... keymap». '
       'После такого ремонта нажмите Preferences → ☰ → Save Preferences, чтобы чистая раскладка '
       'зафиксировалась на диске.')

# ════════════════════════════════════════════════════════════════════════════
# 3. БЫСТРЫЙ СТАРТ
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '3. Быстрый старт')
p(doc, 'Типовой прогон - поставить пивот колесу на ось вращения:')
for b in (
    'Зайдите в режим редактирования колеса, выделите грань на ступице и нажмите D - пивот встанет в центр выделения (Median). Нужно ниже - DD поставит Bottom, ещё одно - DDD, Top.',
    'Нажмите Align Pivot to Edge на панели и ведите курсором по ободу - ось X пивота ложится вдоль ребра; поймали ось колеса - ЛКМ.',
    'Дообработайте наклон, если нужно: Ctrl+Alt+D и поверните ориентацию мышью, Ctrl - шаги по 5 градусам, Shift - тонко.',
    'Позицию подправьте Ctrl+D: пивот скользит по поверхности, Ctrl прижимает к ближайшей вершине, X/Y/Z держат ось.',
    'Перед сдачей прогоните Pivot Audit по конвенции проекта - нарушители покажутся списком, Fix All Flagged исправит пакетно.',
):
    p(doc, b, bullet=True)

# ════════════════════════════════════════════════════════════════════════════
# 4. ORIGIN TO SELECTION
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '4. Origin to Selection (D)')

h2(doc, '4.1 Как работает')
p(doc, 'Выделите что-нибудь и нажмите D - пивот объекта переедет в центр выделения. В режиме '
       'редактирования считается центр выделенных вершин/рёбер/граней; в объектном режиме - '
       'центр выделенных объектов, причём каждый выделенный объект получает свой пивот по своему '
       'выделению. Геометрия при переносе пивота не искажается: вершины переводятся честно, '
       'с компенсацией смещения объекта.')

h2(doc, '4.2 Якоря: D · DD · DDD')
add_table(doc, [
    ('Нажатие', 'Якорь', 'Куда ставится пивот'),
    ('D', 'Median', 'центр выделения'),
    ('DD (быстро, подряд)', 'Bottom', 'XY-медиана выделения, по высоте - нижняя точка'),
    ('DDD (быстро, подряд)', 'Top', 'XY-медиана выделения, по высоте - верхняя точка'),
], [4.4, 3.0, 9.6])
p(doc, 'Четвёртое нажатие подряд (DDDD) возвращает Median, и цикл начинается заново - якорь '
       'ходит по кругу, сколько раз нужно.')

h2(doc, '4.3 Правило паузы')
p(doc, 'Нажатия должны идти подряд, в пределах 0,35 секунды одно за другим - это и есть «DD». '
       'Если сделать паузу дольше, очередное D - уже не следующий якорь, а повторное применение '
       'последнего выбранного. Так же в цикле участвует якорь, выбранный кнопками на панели: '
       'после него очередное DD шагнёт от него к следующему.')

h2(doc, '4.4 Undo')
p(doc, 'Операция коммитится одним шагом истории: один Ctrl+Z возвращает и пивот, и вершины. '
       'После undo сцена сразу отображается корректно - без «улетевших» объектов и повторных '
       'откатов.')

# ════════════════════════════════════════════════════════════════════════════
# 5. DRAG PIVOT
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '5. Drag Pivot (Ctrl+D)')

h2(doc, '5.1 Механика')
p(doc, 'Ctrl+D запускает перетаскивание: пивот следует за курсором, скользя по любой видимой '
       'геометрии сцены (рейкаст). Навигация вьюпорта во время драга продолжает работать. '
       'Во время драга рисуется линия от стартовой позиции и круг-маркер на текущей точке.')

h2(doc, '5.2 Снап к вершинам')
p(doc, 'Зажатый Ctrl включает снап: пивот прилипает к ближайшей вершине в радиусе 24 пикселей '
       'от курсора на экране. Радиус в пикселях, а не в мировых единицах, поэтому работает '
       'одинаково и вблизи, и издалека. Пойманная вершина отмечается жёлтым кругом.')

h2(doc, '5.3 Осевые констеины')
p(doc, 'Клавиши X / Y / Z во время драга ограничивают движение осью: первое нажатие - мировая '
       'ось, повторное той же клавишей - локальная ось объекта, третье - снять ограничение. '
       'Движение по оси идёт экранным скольжением, как у гизмо: ось смотрит почти в камеру - '
       'маркер просто не двигается. Круг-маркер в констеине окрашивается в цвет оси. Снап и '
       'констеин сочетаются: пойманная вершина проецируется на ось.')

h2(doc, '5.4 Подтверждение и отмена')
p(doc, 'ЛКМ или Enter - применить (один undo-шаг). ПКМ или Esc - отмена: до подтверждения сцена '
       'не меняется вовсе, откат возвращает пивот на место. Первый клик по кнопке панели, который '
       'запустил оператор, не засчитывается - защита от случайного мгновенного коммита.')

# ════════════════════════════════════════════════════════════════════════════
# 6. ROTATE PIVOT
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '6. Rotate Pivot (Ctrl+Alt+D)')

h2(doc, '6.1 Механика')
p(doc, 'Оператор объектного режима: вращает ориентацию пивота активного меша, не двигая '
       'геометрию. Зачем: для анимируемых частей - колёс, дверей, крышек - совместить локальную '
       'ось с реальной осью вращения детали, и дальше аниматор крутит её одним каналом. По '
       'умолчанию мышь вращает пивот вокруг оси вида.')

h2(doc, '6.2 Оси и точность')
add_table(doc, [
    ('Клавиша', 'Действие'),
    ('X / Y / Z', 'мировая ось; повторное нажатие - локальная ось объекта; третье - снова ось вида'),
    ('Shift (зажат)', 'тонкое вращение - десятая доля скорости'),
    ('Ctrl (зажат)', 'шаги по 5 градусов'),
    ('Ctrl + Shift', 'шаги по 1 градусу'),
], [4.4, 12.6])

h2(doc, '6.3 Фидбек')
p(doc, 'Во время вращения рисуется живая триада осей пивота и круг-маркер, у курсора выводится '
       'текущий угол с осью («23.4 deg [Z global]»), так же дублируется в статус-бар. Численный '
       'фидбек виден всегда, даже если графику вьюпорта что-то перекрывает. ЛКМ или Enter - '
       'применить, ПКМ или Esc - отмена: до коммита объект не тронут вовсе.')

# ════════════════════════════════════════════════════════════════════════════
# 7. ALIGN PIVOT
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '7. Align Pivot to Normal / to Edge')

h2(doc, '7.1 Align Pivot to Normal')
p(doc, 'Запустите кнопкой панели и ведите курсором по мешу: ось Z пивота следует нормали грани '
       'под курсором, ось Y тянется к мировому верху. Классический случай - дверь: навели '
       'курсором на плоскость полотна - Z смотрит наружу, ЛКМ - применить. Синий круг означает, '
       'что нормаль поймана, жёлтый - грани под курсором нет.')

h2(doc, '7.2 Align Pivot to Edge')
p(doc, 'Тот же жест, но ось X пивота ложится вдоль самого длинного ребра грани под курсором - '
       'детерминированно, без перебора. Нужно другое направление - ведите курсор на соседнюю '
       'грань с нужным ребром. Красный круг - ребро поймано. Для колеса: грань обода даст ось '
       'вдоль обода - это и есть ось вращения.')

h2(doc, '7.3 Липкая цель')
p(doc, 'Цель в обоих режимах «липкая»: увели курсор с меша на панель или в пустоту - последняя '
       'пойманная ориентация сохраняется и ждёт подтверждения. ЛКМ коммитит пойманное даже если '
       'курсор в момент нажатия не над мешем.')

# ════════════════════════════════════════════════════════════════════════════
# 8. PIVOT AUDIT
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '8. Pivot Audit')

h2(doc, '8.1 Зачем')
p(doc, 'Конвенции пивотов плывут: ассеты приходят из разных рук и пакетов, пивоты оказываются '
       'в центре bbox, под землёй или где угодно. Аудит проверяет пивоты выделенных мешей против '
       'выбранного якоря и показывает нарушителей - до того, как ассет поедет дальше по пайплайну.')

h2(doc, '8.2 Правила')
add_table(doc, [
    ('Правило', 'Якорь', 'Допуск'),
    ('Bottom Center', 'низ bbox объекта, XY-центр', 'процент от размера объекта'),
    ('Bounds Center', 'центр bbox объекта', 'процент от размера объекта'),
    ('World Origin', 'мировой ноль (0, 0, 0)', 'метры (по умолчанию 0,1)'),
], [4.6, 6.4, 6.0])
p(doc, 'World Origin - для ассетов, живущих с пивотом в нуле мира (здания, ландшафтные блоки): '
       'геометрия с запасом под землёй здесь - норма, поэтому допуск считается в метрах, а не '
       'в процентах.')

h2(doc, '8.3 Прогон и исправление')
for b in (
    'Фильтр по суффиксу (по умолчанию _geo) отбирает проверяемые объекты - служебные не участвуют.',
    'Audit Selected запускает проверку выделения; нарушители появляются списком с кнопками-селекторами: клик по строке выделяет объект и приближает камеру.',
    'Fix All Flagged пакетно ставит пивот по выбранному якорю; мульти-юзер меши (общая геометрия на несколько объектов) пропускаются с предупреждением - их пивот правится вручную после развязки.',
):
    p(doc, b, bullet=True)

# ════════════════════════════════════════════════════════════════════════════
# 9. ПАНЕЛЬ
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '9. Панель и подсказки')
p(doc, 'Все инструменты живут во вкладке TOCHKA N-панели: кнопки Origin (Median / Bottom / Top), '
       'Drag Pivot, Rotate Pivot, Align to Normal, Align to Edge, кнопка pie-меню со всеми '
       'операторами разом и секция Pivot Audit. Версия расширения выведена в заголовок панели. '
       'Под панелью - свёрнутая суб-панель Info со шпаргалкой клавиш: режим драга (ЛКМ/Enter - '
       'применить, ПКМ/Esc - отмена, Ctrl - снап, X/Y/Z - констеин с циклом глобал-локал-офф), '
       'режим ротации (Shift - тонко, Ctrl - 5 градусов, Ctrl+Shift - 1 градус).')

# ════════════════════════════════════════════════════════════════════════════
# 10. КАК ПОЛЬЗОВАТЬСЯ
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '10. Как пользоваться')

h2(doc, '10.1 Колесо: пивот на ось вращения')
p(doc, 'Align Pivot to Edge по грани обода - ось X легла вдоль обода. Если нужен доводочный '
       'поворот - Ctrl+Alt+D с Ctrl-шагами. Итог: локальная ось колеса совпадает с физической '
       'осью вращения, анимация крутится одним каналом X без компенсационных пустышек.')

h2(doc, '10.2 Дверь: пивот на петли')
p(doc, 'Align Pivot to Normal на плоскость полотна - Z смотрит наружу. Затем Drag Pivot с '
       'констеином по оси - пивот съезжает на кромку петель, Ctrl прилипает к угловой вершине. '
       'Дверь открывается поворотом вокруг своей кромки.')

h2(doc, '10.3 Импорт: привести к конвенции')
p(doc, 'Ассет приехал с пивотом в случайном месте: D с якорем Bottom ставит пивот на землю под '
       'центром объекта. Для парка однотипных объектов быстрее прогнать Audit Selected по правилу '
       'проекта и исправить всё разом.')

h2(doc, '10.4 Проверка перед сдачей')
p(doc, 'Audit Selected с фильтром по суффиксу - последний шаг перед выкладкой ассета: список '
       'нарушителей пуст - пивоты в конвенции; не пуст - Fix All Flagged и перепроверка. '
       'Мульти-юзер меши, если флагнулись, развязать и поправить вручную.')

# ════════════════════════════════════════════════════════════════════════════
# 11. РЕШЕНИЕ ПРОБЛЕМ
# ════════════════════════════════════════════════════════════════════════════

h1(doc, '11. Решение проблем')
add_table(doc, [
    ('Симптом', 'Причина', 'Что делать'),
    ('Клавиша D не запускает оператор', 'бинды раскладки повреждены (старая версия TOCHKA, откат настроек, другой аддон)', 'выключите и включите расширение TOCHKA - сторож восстановит keymap; затем Save Preferences'),
    ('G/R/S перестали работать во вьюпорте', 'повреждены системные keymap Object Mode и Mesh', 'Preferences → Keymap → Restore для Object Mode и Mesh → Save Preferences; TOCHKA 1.0.4+ чинит это сама при включении'),
    ('Снап не цепляет вершину', 'вершина дальше 24 px от курсора на экране', 'приблизьте камеру или ведите курсор точнее - радиус захвата экранный'),
    ('Пивот «улетает» после Ctrl+Z', 'баг версий до 1.0.3 включительно', 'обновите TOCHKA до 1.0.4 или новее'),
    ('Оператор закрылся сразу после запуска кнопкой', 'случился мгновенный коммит от клика по кнопке', 'защита армирования уже вшита: начните движение мыши и подтверждайте ЛКМ'),
], [4.6, 5.4, 7.0])

# ── футер и сохранение ─────────────────────────────────────────────────────

add_footer_pagenum(doc)
doc.save(OUT)
print('saved:', OUT)
